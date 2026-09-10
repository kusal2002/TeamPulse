"""Geometry + text-fit QA for the generated deck (no LibreOffice available here)."""
from pptx import Presentation

EMU = 914400.0
DECK = "TeamPulse-Technical-Presentation.pptx"

prs = Presentation(DECK)
SW = prs.slide_width / EMU
SH = prs.slide_height / EMU
print(f"slide: {SW:.2f} x {SH:.2f} in, {len(prs.slides.__iter__.__self__._sldIdLst)} slides\n")

WIDTH_FACTOR = {"Calibri": 0.47, "Cambria": 0.50, "Courier New": 0.60}
LINE_FACTOR = 1.22

problems = []

for idx, slide in enumerate(prs.slides, start=1):
    boxes = []
    for sh in slide.shapes:
        if sh.left is None or sh.top is None:
            continue
        L, T = sh.left / EMU, sh.top / EMU
        Wd, Ht = (sh.width or 0) / EMU, (sh.height or 0) / EMU
        R, Bm = L + Wd, T + Ht

        if L < -0.01 or T < -0.01 or R > SW + 0.01 or Bm > SH + 0.01:
            problems.append(
                f"S{idx}: OUT OF BOUNDS ({L:.2f},{T:.2f}) {Wd:.2f}x{Ht:.2f} "
                f"right={R:.2f} bottom={Bm:.2f}"
            )

        if Wd < SW - 0.5 and (L < 0.49 or R > SW - 0.49):
            problems.append(f"S{idx}: TIGHT MARGIN left={L:.2f} right={R:.2f}")

        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame
            total_lines, max_sz = 0, 0
            for para in tf.paragraphs:
                text = "".join(r.text for r in para.runs)
                if not text:
                    continue
                sz, face = None, "Calibri"
                for r in para.runs:
                    if r.font.size:
                        sz = r.font.size.pt
                    if r.font.name:
                        face = r.font.name
                    break
                sz = sz or 14
                max_sz = max(max_sz, sz)
                cw = sz * WIDTH_FACTOR.get(face, 0.47) / 72.0
                per_line = max(int(max(Wd - 0.06, 0.2) / cw), 1)
                for seg in text.split("\n"):
                    total_lines += max(1, -(-len(seg) // per_line))
            need = total_lines * max_sz * LINE_FACTOR / 72.0
            if need > Ht + 0.08:
                problems.append(
                    f"S{idx}: TEXT OVERFLOW ~{need:.2f}in vs box {Ht:.2f}in "
                    f"| {tf.text[:55]!r}"
                )
            boxes.append((L, T, R, Bm, tf.text[:32]))

    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            a, b = boxes[i], boxes[j]
            ox = min(a[2], b[2]) - max(a[0], b[0])
            oy = min(a[3], b[3]) - max(a[1], b[1])
            if ox > 0.12 and oy > 0.12:
                problems.append(
                    f"S{idx}: TEXT OVERLAP {ox:.2f}x{oy:.2f}in | {a[4]!r} vs {b[4]!r}"
                )

if problems:
    print(f"{len(problems)} issue(s):\n")
    for p in problems:
        print(" -", p)
else:
    print("No geometry or text-fit issues found.")
